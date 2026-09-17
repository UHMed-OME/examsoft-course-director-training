#!/usr/bin/env python3
"""build_deck.py — renders a 16:9 slide deck from the same content/ JSON the
website is built from.

    python build/build_deck.py

Output: dist/ExamSoft-for-Course-Directors.pptx

DESIGN NOTE, please read before adding slides.

This deck is for a live session. It is not the website on slides. The website
is the reference and it can be dense, because someone reads it alone with time
to spare. A deck competes with the presenter for the audience's attention, so
every word on screen is a word nobody is listening to.

So the rules here are:

  * One idea per slide.
  * Roughly 25 words on screen, and never more than four bullets.
  * Detail belongs in the speaker notes, which is where the presenter needs it
    anyway. The notes are not a dumping ground; they are the script.
  * Tables get trimmed to the rows that change behavior. The full table lives
    on the website, and the deck points there.

The slide plan below is explicit rather than generated per section. Generating
one slide per section is what produced the bloated first version.

Requires python-pptx (pip install python-pptx). Nothing else.
"""

import json
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("python-pptx is not installed. Run: pip install python-pptx")

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "content"
DIST = ROOT / "dist"
IMGCACHE = DIST / "_img"


def optimized(path, max_w=1100, quality=84):
    """Downscale an image for slide use.

    The originals in images/ stay full resolution for the website. A 1431px
    photograph displayed across seven inches of a slide is wasted bytes, and it
    slows both the Drive upload and the Google Slides conversion. Cached on
    mtime so repeat builds are cheap.
    """
    from PIL import Image

    path = Path(path)
    IMGCACHE.mkdir(parents=True, exist_ok=True)
    out = IMGCACHE / f"{path.stem}.jpg"
    if out.exists() and out.stat().st_mtime >= path.stat().st_mtime:
        return out

    with Image.open(path) as im:
        im = im.convert("RGB")
        if im.width > max_w:
            im = im.resize((max_w, round(im.height * max_w / im.width)),
                           Image.LANCZOS)
        im.save(out, "JPEG", quality=quality, optimize=True, progressive=True)
    return out


# --- Palette -----------------------------------------------------------------
# Mirrors assets/css/tokens.css. Change these when you change the tokens.
BRAND = RGBColor(0x02, 0x47, 0x31)
BRAND_LT = RGBColor(0x6F, 0xC3, 0x9F)
SOFT = RGBColor(0xEE, 0xF5, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x14, 0x20, 0x1B)
MUTED = RGBColor(0x4A, 0x5A, 0x52)
FAINT = RGBColor(0x66, 0x75, 0x6D)

CRITICAL = RGBColor(0xA4, 0x20, 0x1A)
WARN = RGBColor(0x8A, 0x5A, 0x00)
NOTE = RGBColor(0x1C, 0x5B, 0x96)

TONE_COLOR = {"critical": CRITICAL, "warn": WARN, "note": NOTE, "jabsom": BRAND}
TONE_LABEL = {"critical": "REMEMBER THIS", "warn": "WATCH OUT",
              "note": "NOTE", "jabsom": "JABSOM RULE"}

FONT = "Segoe UI"

SW = Inches(13.333)
SH = Inches(7.5)
M = Inches(0.85)
CONTENT_W = SW - (2 * M)


# --- Text helpers ------------------------------------------------------------
def plain(text):
    """Strip the site's inline markup. Slides get flat text."""
    if text is None:
        return ""
    s = str(text)
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = s.replace("`", "")
    return re.sub(r"\s+", " ", s).strip()


def clip(text, limit):
    """Trim to a sentence boundary where possible, so bullets do not dangle."""
    t = plain(text)
    if len(t) <= limit:
        return t
    cut = t[:limit]
    for sep in (". ", "; ", ", "):
        i = cut.rfind(sep)
        if i > limit * 0.55:
            return cut[: i + 1].rstrip(" ;,")
    return cut.rsplit(" ", 1)[0].rstrip(" ,;:") + "..."


def textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size, color, bold=False, space_after=6,
         first=False, align=PP_ALIGN.LEFT, spacing=None, caps=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if spacing:
        p.line_spacing = spacing
    run = p.add_run()
    run.text = text.upper() if caps else text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return p


def rect(slide, left, top, width, height, fill):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def set_alt(shape, description):
    """python-pptx has no alt-text API, so write the descr attribute directly.
    PowerPoint and Google Slides both read it."""
    shape._element._nvXxPr.cNvPr.set("descr", description)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def notes(slide, text):
    """The presenter's script. Everything trimmed off a slide lands here."""
    slide.notes_slide.notes_text_frame.text = plain(text)


def header(slide, eyebrow, title):
    tf = textbox(slide, M, Inches(0.62), CONTENT_W, Inches(1.5))
    para(tf, plain(eyebrow), 12, BRAND, bold=True, first=True,
         caps=True, space_after=6)
    para(tf, plain(title), 32, INK, bold=True, space_after=0, spacing=1.0)
    rect(slide, M, Inches(1.82), Inches(1.5), Emu(22860), BRAND)


# --- Slide builders ----------------------------------------------------------
def slide_title(prs, site):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, BRAND)
    rect(s, 0, 0, Inches(0.22), SH, BRAND_LT)

    tf = textbox(s, M, Inches(1.6), CONTENT_W, Inches(3.2))
    para(tf, f"{site['institution']} · {site['office']}", 13, BRAND_LT,
         bold=True, first=True, caps=True, space_after=16)
    para(tf, site["title"], 46, WHITE, bold=True, space_after=12, spacing=1.0)
    para(tf, "Onboarding for new pre-clerkship course directors", 19, SOFT,
         space_after=0, spacing=1.2)

    tf2 = textbox(s, M, Inches(5.6), CONTENT_W, Inches(1.2))
    para(tf2, "INSTITUTION ID", 11, BRAND_LT, bold=True, first=True, space_after=2)
    para(tf2, site["portal"]["institutionId"], 20, WHITE, bold=True, space_after=10)
    para(tf2, f"{site['portal']['edition']} portal · {site['academicYear']}",
         12, SOFT, space_after=0)

    notes(s, "Faculty counterpart to the student Examplify orientation deck. "
             "Generated from the training repository, so do not edit slides by "
             "hand. Point people at the website for anything not on a slide.")
    return s


def slide_agenda(prs, site):
    """Four themes, not a table of contents. Thirteen line items is a wall."""
    s = blank(prs)
    header(s, "Today", "Four things")

    items = [
        ("Get set up", "Sign in, and set your fonts before you write anything."),
        ("Write questions", "The bank, the tagging rule, and the two question types worth learning."),
        ("Get it delivered", "The 7 day deadline and what OME does with your draft."),
        ("Read the results", "What to look at after students upload."),
    ]

    gap = Inches(0.3)
    col_w = (CONTENT_W - gap * 3) / 4
    for i, (title, detail) in enumerate(items):
        left = int(M + i * (col_w + gap))
        rect(s, left, Inches(2.5), int(col_w), Inches(0.05), BRAND)
        tf = textbox(s, left, Inches(2.8), int(col_w), Inches(2.6))
        para(tf, f"{i + 1:02d}", 13, BRAND_LT, bold=True, first=True, space_after=8)
        para(tf, title, 19, INK, bold=True, space_after=10, spacing=1.1)
        para(tf, detail, 12, MUTED, space_after=0, spacing=1.25)

    notes(s, "Keep this short. The detail is on the website. This slide only "
             "sets expectations for the session.")
    return s


def slide_big(prs, eyebrow, headline, tone, sub=None, note=None):
    """One statement, nothing else. Used for the rules that actually matter."""
    s = blank(prs)
    color = TONE_COLOR.get(tone, NOTE)
    header(s, eyebrow, "")

    tf = textbox(s, M, Inches(2.5), CONTENT_W * 0.88, Inches(3.0))
    para(tf, TONE_LABEL.get(tone, "NOTE"), 12, color, bold=True, first=True,
         space_after=16)
    para(tf, plain(headline), 40, INK, bold=True, space_after=20, spacing=1.05)
    if sub:
        para(tf, clip(sub, 190), 17, MUTED, space_after=0, spacing=1.3)

    rect(s, M, Inches(2.5), Inches(0.07), Inches(2.4), color)
    if note:
        notes(s, note)
    return s


def slide_bullets(prs, eyebrow, title, items, note=None):
    s = blank(prs)
    header(s, eyebrow, title)

    tf = textbox(s, M, Inches(2.5), CONTENT_W * 0.86, Inches(4.2))
    for i, item in enumerate(items[:4]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(18)
        p.line_spacing = 1.25
        dot = p.add_run()
        dot.text = "—  "
        dot.font.size = Pt(17)
        dot.font.bold = True
        dot.font.color.rgb = BRAND
        dot.font.name = FONT
        run = p.add_run()
        run.text = clip(item, 95)
        run.font.size = Pt(17)
        run.font.color.rgb = INK
        run.font.name = FONT
    if note:
        notes(s, note)
    return s


def slide_steps(prs, eyebrow, title, steps, note=None):
    s = blank(prs)
    header(s, eyebrow, title)

    gap = Inches(0.4)
    col_w = (CONTENT_W - gap * (len(steps) - 1)) / len(steps)
    for i, (label, detail) in enumerate(steps):
        left = int(M + i * (col_w + gap))
        circle = rect(s, left, Inches(2.5), Inches(0.5), Inches(0.5), BRAND)
        ctf = circle.text_frame
        ctf.margin_left = ctf.margin_right = 0
        ctf.margin_top = ctf.margin_bottom = 0
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(ctf, str(i + 1), 16, WHITE, bold=True, first=True,
             align=PP_ALIGN.CENTER, space_after=0)

        tf = textbox(s, left, Inches(3.25), int(col_w), Inches(2.6))
        para(tf, label, 20, INK, bold=True, first=True, space_after=10, spacing=1.1)
        para(tf, clip(detail, 140), 13, MUTED, space_after=0, spacing=1.3)
    if note:
        notes(s, note)
    return s


def slide_timeline(prs, eyebrow, title, items, note=None):
    s = blank(prs)
    header(s, eyebrow, title)

    n = len(items)
    gap = Inches(0.18)
    col_w = (CONTENT_W - gap * (n - 1)) / n
    top = Inches(2.6)

    for i, item in enumerate(items):
        left = int(M + i * (col_w + gap))
        rect(s, left, top, int(col_w), Inches(0.055), BRAND)
        tf = textbox(s, left, top + Inches(0.28), int(col_w), Inches(3.4))
        para(tf, plain(item["when"]), 11, BRAND, bold=True, first=True,
             caps=True, space_after=8, spacing=1.1)
        para(tf, plain(item["what"]), 17, INK, bold=True, space_after=7, spacing=1.1)
        para(tf, plain(item["who"]), 10, FAINT, bold=True, caps=True, space_after=0)
    if note:
        notes(s, note)
    return s


def slide_image(prs, eyebrow, title, image, alt, caption, note=None):
    s = blank(prs)
    header(s, eyebrow, title)

    from PIL import Image
    src = optimized(image)
    with Image.open(src) as im:
        iw, ih = im.size

    scale = min((CONTENT_W * 0.66) / iw, Inches(4.0) / ih)
    w, h = int(iw * scale), int(ih * scale)

    pic = s.shapes.add_picture(str(src), M, Inches(2.4), width=w, height=h)
    set_alt(pic, alt)

    tx = M + w + Inches(0.55)
    tf = textbox(s, tx, Inches(2.4), SW - tx - M, Inches(4.0))
    para(tf, clip(caption, 150), 16, INK, first=True, space_after=0, spacing=1.3)

    notes(s, (note + " " if note else "") + "Alt text: " + alt)
    return s


def slide_image_pair(prs, eyebrow, title, items, note=None):
    s = blank(prs)
    header(s, eyebrow, title)

    from PIL import Image
    gap = Inches(0.45)
    col_w = (CONTENT_W - gap) / 2
    top = Inches(2.45)

    for i, item in enumerate(items):
        path = optimized(ROOT / item["src"], max_w=900)
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(col_w / iw, Inches(3.4) / ih)
        w, h = int(iw * scale), int(ih * scale)
        left = int(M + i * (col_w + gap) + (col_w - w) / 2)
        pic = s.shapes.add_picture(str(path), left, top, width=w, height=h)
        set_alt(pic, item["alt"])

        tf = textbox(s, int(M + i * (col_w + gap)), top + h + Inches(0.2),
                     int(col_w), Inches(0.5))
        para(tf, plain(item["label"]), 14, INK, bold=True, first=True,
             space_after=0, align=PP_ALIGN.CENTER)

    notes(s, (note + " " if note else "") +
          " | ".join(i["alt"] for i in items))
    return s


def slide_table(prs, eyebrow, title, cols, rows, widths=None, note=None):
    s = blank(prs)
    header(s, eyebrow, title)

    top = Inches(2.4)
    shape = s.shapes.add_table(len(rows) + 1, len(cols), M, top,
                               int(CONTENT_W), int(SH - top - Inches(0.7)))
    table = shape.table

    if widths:
        total = sum(widths)
        for i, frac in enumerate(widths):
            table.columns[i].width = int(CONTENT_W * frac / total)

    for c, label in enumerate(cols):
        cell = table.cell(0, c)
        cell.text = plain(label)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRAND
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.13)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = FONT

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = plain(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else SOFT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.13)
            cell.margin_top = cell.margin_bottom = Inches(0.06)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(13)
            p.font.bold = c == 0
            p.font.color.rgb = INK if c == 0 else MUTED
            p.font.name = FONT
    if note:
        notes(s, note)
    return s


def slide_contacts(prs, site):
    s = blank(prs)
    header(s, "Support", "Who to ask")

    gap = Inches(0.4)
    n = len(site["contacts"])
    col_w = (CONTENT_W - gap * (n - 1)) / n

    for i, c in enumerate(site["contacts"]):
        left = int(M + i * (col_w + gap))
        rect(s, left, Inches(2.5), int(col_w), Inches(2.4), SOFT)
        tf = textbox(s, left + Inches(0.4), Inches(2.85),
                     int(col_w - Inches(0.8)), Inches(1.9))
        para(tf, plain(c["role"]), 10, BRAND, bold=True, first=True,
             caps=True, space_after=10)
        para(tf, plain(c["name"]), 22, INK, bold=True, space_after=4)
        para(tf, c["email"], 14, BRAND, bold=True, space_after=0)

    tf = textbox(s, M, Inches(5.5), CONTENT_W * 0.85, Inches(1.0))
    para(tf, "OME files vendor tickets, not you", 17, INK, bold=True,
         first=True, space_after=6)
    para(tf, "Send anything that looks like an ExamSoft defect to Jesse.",
         14, MUTED, space_after=0)

    notes(s, site["vendorTicketPolicy"])
    return s


def slide_closing(prs, site):
    s = blank(prs)
    rect(s, 0, 0, SW, SH, BRAND)
    rect(s, 0, 0, Inches(0.22), SH, BRAND_LT)
    tf = textbox(s, M, Inches(2.7), CONTENT_W, Inches(2.4))
    para(tf, "Mahalo", 54, WHITE, bold=True, first=True, space_after=18)
    para(tf, f"Questions any time at {site['contacts'][0]['email']}",
         19, SOFT, space_after=16)
    para(tf, "Everything in this deck, in full, is on the course director site.",
         13, BRAND_LT, space_after=0)
    notes(s, "Leave this up during questions.")
    return s


# --- Assembly ----------------------------------------------------------------
def find(section, block_type, **match):
    for b in section["blocks"]:
        if b.get("type") != block_type:
            continue
        if all(b.get(k) == v for k, v in match.items()):
            return b
    return None


def main():
    site = json.loads((CONTENT / "site.json").read_text(encoding="utf-8"))
    home = json.loads((CONTENT / "home.json").read_text(encoding="utf-8"))
    sections = {
        json.loads(f.read_text(encoding="utf-8"))["id"]:
        json.loads(f.read_text(encoding="utf-8"))
        for f in sorted((CONTENT / "sections").glob("*.json"))
    }

    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # 1. Frame the session
    slide_title(prs, site)
    slide_agenda(prs, site)

    who = find(home, "table")
    slide_table(prs, "Your role", "Who does what",
                who["cols"], who["rows"], widths=[1, 1],
                note="Course directors often assume they post the exam "
                     "themselves. They do not. Make this split explicit now and "
                     "it prevents most of the confusion later.")

    # 2. The deadline, on its own
    slide_big(prs, "The deadline",
              "Your draft is due 7 business days before the exam",
              "critical",
              "Every other date is scheduled backwards from that one.",
              note="This is the single most important slide in the deck. "
                   "Proofing, correction, re-preview, printing and proctor prep "
                   "run in sequence and need the full window. A late draft "
                   "compresses proofing. It does not move the exam date. The "
                   "2020 onboarding document put it best: it is more than just "
                   "clicking a button.")

    slide_timeline(prs, "The deadline", "How an exam gets delivered",
                   find(home, "timeline")["items"],
                   note="Walk left to right. Detail for each stage is on the "
                        "website under The JABSOM exam cycle. Stress that the "
                        "middle stage is where staff find the errors.")

    # 3. Get set up
    slide_steps(prs, "Week one", "Do these three things",
                [("Sign in", "Institution ID is hawaiimed. Uses your UH login, so there is no separate password."),
                 ("Set fonts to Arial 12", "My Preferences, Question Creation tab. Before you write anything."),
                 ("Open last year's exam", "You will build this year's version from it.")],
                note="If sign-in fails with 'No associated ExamSoft account "
                     "found', that is an External ID mismatch on the admin "
                     "side. Email Jesse. An incognito window will not help, and "
                     "people waste time trying.")

    pref = sections["preferences"]
    slide_image(prs, "Week one", "Set both font fields to Arial 12",
                ROOT / find(pref, "figure")["src"],
                find(pref, "figure")["alt"],
                "My Preferences, Question Creation tab. Set both, then Save.",
                note="Nine years of items from different course directors end "
                     "up in one assessment. Arial 12 everywhere is what keeps "
                     "that looking consistent. Two fields and a Save button.")

    # 4. Writing questions
    slide_big(prs, "Writing questions",
              "Every question needs at least one category tag",
              "jabsom",
              "An untagged question disappears from category reporting.",
              note="The software will not enforce this. JABSOM does. Category "
                   "reporting is how the school looks at performance by topic "
                   "across a course and across years, and an untagged item is "
                   "invisible to it. Tag while writing, or bulk-tag later from "
                   "the questions table. Apply tags, but do not restructure the "
                   "tree: it is shared across all historical data.")

    qt = sections["question-types"]
    slide_table(prs, "Writing questions", "Question types at a glance",
                ["Type", "Use it?"],
                [[r[0], r[1]] for r in find(qt, "table")["rows"]],
                widths=[1.4, 1],
                note="Fill in the Blank and Matching both work. OME discourages "
                     "them because they generate more scoring disputes than "
                     "they are worth. The full table with notes is on the "
                     "website.")

    slide_image(prs, "Question types", "Hot Spot",
                ROOT / find(qt, "figure")["src"],
                find(qt, "figure")["alt"],
                "The student taps the image and a pin drops.",
                note="Best item type we have for anatomy, neuroanatomy and "
                     "clinical skills, or anywhere the answer is a location "
                     "rather than a word. Rectangle or polygon. Students can "
                     "pinch to zoom, so small structures still work.")

    pair = find(qt, "figurepair")
    slide_image_pair(prs, "Question types", "Drag and Drop", pair["items"],
                     note="Student drags choices from the left box to the "
                          "right. Keep choices short.")

    slide_big(prs, "Question types",
              "Drag and Drop does not score the order",
              "critical",
              "And with partial credit on, a wrong answer subtracts points.",
              note="Two things people get wrong. First, sequence does not "
                   "matter, so if you need order scored use the Ordering type. "
                   "Second, partial credit here is symmetric: a correct answer "
                   "adds points and an incorrect one subtracts the same amount, "
                   "floored at zero. Say so in the stem or students guess.")

    att = sections["attachments"]
    slide_table(prs, "Attachments", "The limits that will stop you",
                ["Type", "Max size"],
                [[r[0], r[2]] for r in find(att, "table")["rows"]],
                widths=[1.3, 1],
                note="200 MB total per assessment, covering every attachment "
                     "and inline image together. Image-heavy exams reach it. "
                     "Also: assessment level takes exactly one file, question "
                     "level takes several.")

    # 5. Delivery
    bld = sections["building"]
    slide_big(prs, "Building an assessment",
              "Copy last year's exam, then edit it",
              "jabsom",
              "With nine years of items on the system, almost nobody starts from scratch.",
              note="Duplicate, do not repost. A repost folds this year's "
                   "results into last year's reports and ruins year-over-year "
                   "comparison. Duplicating gives the exam its own Exam ID and "
                   "its own reporting. The comparison table is on the website.")

    slide_big(prs, "Posting", "One download and the questions lock",
              "critical",
              "After that, an error becomes a scoring adjustment instead of an edit.",
              note="OME posts your exam. Two things to know anyway. Max "
                   "Downloads is always 1, which is a standing JABSOM rule. And "
                   "once the exam is posted and one student has downloaded it, "
                   "questions and answer keys freeze. That is why proofing "
                   "happens before the download window opens. The password is "
                   "never emailed to students; the proctor announces it.")

    slide_bullets(prs, "Exam day", "What happens in the room",
                  ["Examplify locks the iPad. No screenshots, no copy and paste, no internet.",
                   "Students cannot open the exam until the proctor announces the password.",
                   "Too many button combinations locks a student out. The proctor holds the resume code.",
                   "OME prints two hard copies as a backup."],
                  note="Students are told to charge the iPad fully and restart "
                       "it once beforehand, which clears most problems. If "
                       "something fails mid-exam they tell the proctor rather "
                       "than troubleshooting. Resume codes live on the proctor "
                       "PowerPoint OME prepares, never in email.")

    # 6. After
    slide_bullets(prs, "After the exam", "Results arrive in minutes",
                  ["The assessment overview shows how the class did as soon as students upload.",
                   "Item analysis finds questions the whole class missed.",
                   "Category performance is what the tags were for.",
                   "Talk to OME before adjusting scores. It changes the record for everyone."],
                  note="Point out the annual loop: this year's item analysis is "
                       "what tells you which questions to cut when you duplicate "
                       "the exam next year.")

    slide_contacts(prs, site)
    slide_closing(prs, site)

    DIST.mkdir(exist_ok=True)
    out = DIST / "ExamSoft-for-Course-Directors.pptx"
    prs.save(out)

    count = len(prs.slides._sldIdLst)
    print(f"Built {out.relative_to(ROOT)} with {count} slides")
    if count > 22:
        print("WARNING: over 22 slides. Re-read the design note at the top of "
              "this file before adding more.")


if __name__ == "__main__":
    main()
